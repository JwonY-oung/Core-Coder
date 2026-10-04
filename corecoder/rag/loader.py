"""Document loading: turn files on disk into plain text.

The rest of the RAG pipeline only ever sees strings, so every format quirk is
absorbed here. Two tiers:

  * always available - .txt / .md / .rst / .markdown / .csv / .json / .py ...
    anything UTF-8 readable. No dependency, works offline.
  * optional - .pdf (pypdf) and .docx (python-docx). Imported lazily inside the
    loader so a missing extra never breaks the text path; when the extra isn't
    installed we raise :class:`LoaderError` with the exact pip command, and the
    caller skips the file instead of dying.

Each loader returns a list of :class:`LoadedDoc` rather than one blob, because
PDFs and Word files have a natural page/paragraph boundary worth keeping as
metadata: it ends up in the citation the customer-service bot shows the user.
"""

from __future__ import annotations

import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

# read without an extra -> the file is almost certainly binary junk, not a doc
MAX_BYTES = 20 * 1024 * 1024

# text-ish extensions we accept with the plain reader
TEXT_EXTS = {
    ".txt", ".md", ".markdown", ".rst", ".text", ".adoc",
    ".csv", ".tsv", ".json", ".yaml", ".yml", ".toml", ".ini", ".cfg",
    ".py", ".js", ".ts", ".java", ".go", ".rs", ".c", ".h", ".cpp", ".sh",
    ".html", ".htm", ".xml", ".log", ".faq",
}
PDF_EXTS = {".pdf"}
DOCX_EXTS = {".docx"}

# directories that are never knowledge, only noise
SKIP_DIRS = {".git", ".hg", ".svn", "node_modules", "__pycache__", ".venv", "venv",
             ".tox", "dist", "build", ".idea", ".vscode", ".pytest_cache", ".mypy_cache"}


class LoaderError(RuntimeError):
    """A file could not be turned into text (unreadable, binary, missing extra)."""


@dataclass
class LoadedDoc:
    """One unit of loaded text plus where it came from."""

    text: str
    source: str
    meta: dict = field(default_factory=dict)

    @property
    def is_empty(self) -> bool:
        return not self.text.strip()


# ---------------------------------------------------------------------------
# per-format readers
# ---------------------------------------------------------------------------

def _read_text(path: Path) -> str:
    """UTF-8 with replacement, so a stray latin-1 byte can't kill the ingest."""
    if path.stat().st_size > MAX_BYTES:
        raise LoaderError(f"{path.name}: larger than {MAX_BYTES // (1024 * 1024)}MB, skipped")
    try:
        return path.read_text(encoding="utf-8", errors="replace")
    except (OSError, UnicodeError) as e:
        raise LoaderError(f"{path.name}: {e}") from e


def _strip_markdown(text: str) -> str:
    """Drop the markup that adds nothing to retrieval.

    Kept deliberately shallow: we only remove link syntax and emphasis marks,
    because those actively hurt tokenisation ("[退款](url)" tokenises the URL
    into the chunk). Headings keep their text - a heading is often the single
    best retrieval signal in a FAQ.
    """
    import re

    text = re.sub(r"!\[([^\]]*)\]\([^)]*\)", r"\1", text)   # images -> alt text
    text = re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", text)     # links -> label
    text = re.sub(r"^\s{0,3}#{1,6}\s*", "", text, flags=re.M)  # heading marks
    text = re.sub(r"(\*\*|__|\*|_|`)", "", text)             # emphasis / code ticks
    return text


def _read_json_doc(path: Path) -> str:
    """JSON knowledge (FAQ exports, structured articles) flattened to lines.

    A raw JSON dump tokenises terribly - braces and quotes dominate. We walk it
    and emit ``dotted.path: value`` lines, which both BM25 and the hashing
    embedder can actually match against.
    """
    raw = _read_text(path)
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        return raw  # not valid JSON after all; index it as plain text

    lines: list[str] = []

    def walk(node, prefix: str) -> None:
        if isinstance(node, dict):
            for k, v in node.items():
                walk(v, f"{prefix}.{k}" if prefix else str(k))
        elif isinstance(node, list):
            for i, v in enumerate(node):
                walk(v, f"{prefix}[{i}]")
        else:
            if node is None or node == "":
                return
            lines.append(f"{prefix}: {node}" if prefix else str(node))

    walk(data, "")
    return "\n".join(lines) or raw


def _read_pdf(path: Path) -> list[LoadedDoc]:
    """One LoadedDoc per page, so citations can say "manual.pdf p.12"."""
    try:
        from pypdf import PdfReader
    except ImportError as e:  # pragma: no cover - depends on the optional extra
        raise LoaderError(f"{path.name}: PDF support needs pypdf (pip install 'corecoder[rag]')") from e
    try:
        reader = PdfReader(str(path))
    except Exception as e:  # noqa: BLE001 - pypdf raises a zoo of parse errors
        raise LoaderError(f"{path.name}: unreadable PDF ({e})") from e

    docs: list[LoadedDoc] = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:  # noqa: BLE001 - a broken page shouldn't lose the rest
            text = ""
        if text.strip():
            docs.append(LoadedDoc(text=text, source=f"{path.name}#p{i}", meta={"page": i}))
    if not docs:
        raise LoaderError(f"{path.name}: no extractable text (scanned PDF? needs OCR)")
    return docs


def _read_docx(path: Path) -> LoadedDoc:
    try:
        import docx  # python-docx
    except ImportError as e:  # pragma: no cover
        raise LoaderError(f"{path.name}: Word support needs python-docx (pip install 'corecoder[rag]')") from e
    try:
        document = docx.Document(str(path))
    except Exception as e:  # noqa: BLE001
        raise LoaderError(f"{path.name}: unreadable .docx ({e})") from e

    parts = [p.text for p in document.paragraphs if p.text.strip()]
    # tables carry half the content of a typical spec sheet; flatten row by row
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n\n".join(parts)
    if not text.strip():
        raise LoaderError(f"{path.name}: empty document")
    return LoadedDoc(text=text, source=path.name, meta={"paragraphs": len(parts)})


# ---------------------------------------------------------------------------
# public API
# ---------------------------------------------------------------------------

def supported(path: str | Path) -> bool:
    """True if :func:`load_file` knows how to read this extension."""
    ext = Path(path).suffix.lower()
    return ext in TEXT_EXTS or ext in PDF_EXTS or ext in DOCX_EXTS


def load_file(path: str | Path, *, strip_markdown: bool = True) -> list[LoadedDoc]:
    """Read one file into LoadedDocs. Raises LoaderError if it can't."""
    p = Path(path).expanduser()
    if not p.is_file():
        raise LoaderError(f"not a file: {p}")
    ext = p.suffix.lower()

    if ext in PDF_EXTS:
        return _read_pdf(p)
    if ext in DOCX_EXTS:
        return [_read_docx(p)]
    if ext not in TEXT_EXTS:
        raise LoaderError(f"unsupported extension: {ext or p.name}")

    text = _read_json_doc(p) if ext == ".json" else _read_text(p)
    if strip_markdown and ext in (".md", ".markdown"):
        text = _strip_markdown(text)
    if not text.strip():
        raise LoaderError(f"{p.name}: empty")
    return [LoadedDoc(text=text, source=str(p), meta={"ext": ext})]


def iter_files(
    directory: str | Path,
    *,
    recursive: bool = True,
    glob: str = "**/*",
) -> list[Path]:
    """Supported files under a directory, sorted, hidden/skip-dirs excluded."""
    root = Path(directory).expanduser()
    if not root.is_dir():
        raise NotADirectoryError(f"not a directory: {root}")
    it = root.glob(glob if recursive else "*")
    out: list[Path] = []
    for p in sorted(it):
        if not p.is_file() or p.name.startswith("."):
            continue
        if any(part in SKIP_DIRS or part.startswith(".") for part in p.relative_to(root).parts[:-1]):
            continue
        if supported(p):
            out.append(p)
    return out


def load_files(
    paths: Iterable[str | Path],
    *,
    on_error=None,
    **kwargs,
) -> list[LoadedDoc]:
    """Load an explicit list of files (a glob already did the filtering).

    ``on_error(path, message)`` is called per skip instead of raising, so a
    caller ingesting ``docs/**/*.pdf`` keeps the 199 readable files when one is
    a scan.
    """
    docs: list[LoadedDoc] = []
    for raw in paths:
        p = Path(raw).expanduser()
        try:
            docs.extend(load_file(p, **kwargs))
        except (LoaderError, OSError) as e:
            if on_error:
                on_error(p, str(e))
    return docs


def load_dir(
    directory: str | Path,
    *,
    recursive: bool = True,
    on_error=None,
    glob: str = "**/*",
) -> list[LoadedDoc]:
    """Load every supported file under a directory.

    Failures are collected, not raised: one scanned PDF shouldn't abort the
    ingest of 200 good markdown files. ``on_error(path, message)`` gets each
    skip so the CLI can report it.
    """
    return load_files(
        iter_files(directory, recursive=recursive, glob=glob), on_error=on_error
    )


__all__ = [
    "DOCX_EXTS",
    "LoaderError",
    "LoadedDoc",
    "PDF_EXTS",
    "TEXT_EXTS",
    "iter_files",
    "load_dir",
    "load_file",
    "load_files",
    "supported",
]
